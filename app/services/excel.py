from datetime import date
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.services.ledger import (
    list_invoice_records,
    list_order_records
)


def parse_excel_date(value):
    if not value:
        return None
    try:
        return date.fromisoformat(value)
    except (TypeError, ValueError):
        return value


def style_export_sheet(
    worksheet,
    column_widths: list[int]
) -> None:
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    worksheet.sheet_view.showGridLines = False
    header_fill = PatternFill(
        fill_type="solid",
        fgColor="2563EB"
    )

    for cell in worksheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True
        )
    worksheet.row_dimensions[1].height = 28

    for column_index, width in enumerate(
        column_widths,
        start=1
    ):
        column_letter = get_column_letter(column_index)
        worksheet.column_dimensions[column_letter].width = width

    for row in worksheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(
                vertical="top",
                wrap_text=True
            )


def build_ledger_workbook() -> BytesIO:
    orders = list_order_records()
    invoices = list_invoice_records()
    workbook = Workbook()

    order_sheet = workbook.active
    order_sheet.title = "订单发票总表"
    order_sheet.append([
        "订单类型",
        "订单号",
        "合同号",
        "往来单位",
        "订单日期",
        "交付日期",
        "订单金额",
        "预付款金额",
        "累计收付款",
        "剩余金额",
        "到期日",
        "收付款状态",
        "开票状态",
        "关联发票",
        "结算方式",
        "付款账期",
        "负责人",
        "复核状态"
    ])

    for order in orders:
        order_sheet.append([
            order["order_type"],
            order["order_number"],
            order["contract_number"],
            order["counterparty_name"],
            parse_excel_date(order["order_date"]),
            parse_excel_date(order["delivery_date"]),
            float(order["order_amount"] or 0),
            float(order["prepayment_amount"] or 0),
            float(order["paid_amount"] or 0),
            None,
            parse_excel_date(order["due_date"]),
            order["payment_status"],
            order["invoice_status"],
            "、".join(order["invoice_numbers"] or []),
            order["settlement_method"],
            order["payment_terms"],
            order["owner"],
            order["review_status"]
        ])
        row_number = order_sheet.max_row
        order_sheet.cell(
            row=row_number,
            column=10,
            value=f"=MAX(0,G{row_number}-I{row_number})"
        )

    style_export_sheet(
        order_sheet,
        [
            12, 20, 20, 28, 13, 13,
            15, 15, 15, 15, 13, 14,
            14, 30, 20, 32, 14, 14
        ]
    )
    for row_number in range(2, order_sheet.max_row + 1):
        for column in (7, 8, 9, 10):
            order_sheet.cell(
                row=row_number,
                column=column
            ).number_format = "#,##0.00"
        for column in (5, 6, 11):
            order_sheet.cell(
                row=row_number,
                column=column
            ).number_format = "yyyy-mm-dd"

    invoice_sheet = workbook.create_sheet("发票明细")
    invoice_sheet.append([
        "发票类型",
        "发票号码",
        "开票日期",
        "购买方",
        "销售方",
        "未税金额",
        "税额",
        "价税合计",
        "预付款金额",
        "累计收付款",
        "剩余金额",
        "到期日",
        "收付款状态",
        "关联订单",
        "负责人",
        "复核状态",
        "来源文件"
    ])

    for invoice in invoices:
        invoice_sheet.append([
            invoice["document_type"],
            invoice["invoice_number"],
            parse_excel_date(invoice["invoice_date"]),
            invoice["buyer_name"],
            invoice["seller_name"],
            float(invoice["amount_excl_tax"] or 0),
            float(invoice["tax_amount"] or 0),
            float(invoice["amount_incl_tax"] or 0),
            float(invoice["prepayment_amount"] or 0),
            float(invoice["paid_amount"] or 0),
            None,
            parse_excel_date(invoice["due_date"]),
            invoice["payment_status"],
            "、".join(invoice["order_numbers"] or []),
            invoice["owner"],
            invoice["review_status"],
            invoice["source_filename"]
        ])
        row_number = invoice_sheet.max_row
        invoice_sheet.cell(
            row=row_number,
            column=11,
            value=f"=MAX(0,H{row_number}-J{row_number})"
        )

    style_export_sheet(
        invoice_sheet,
        [
            12, 24, 13, 30, 30, 15,
            15, 15, 15, 15, 15, 13,
            14, 32, 14, 14, 30
        ]
    )
    for row_number in range(2, invoice_sheet.max_row + 1):
        for column in (6, 7, 8, 9, 10, 11):
            invoice_sheet.cell(
                row=row_number,
                column=column
            ).number_format = "#,##0.00"
        for column in (3, 12):
            invoice_sheet.cell(
                row=row_number,
                column=column
            ).number_format = "yyyy-mm-dd"

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output
