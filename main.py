import os
import json
import sqlite3
import binascii
import secrets
from contextlib import asynccontextmanager, closing
from apscheduler.schedulers.asyncio import AsyncIOScheduler
from pathlib import Path
import httpx
from dotenv import load_dotenv
from fastapi import (
    FastAPI,
    File,
    Form,
    HTTPException,
    Request,
    Response,
    UploadFile
)
from typing import Any, Literal
from pydantic import BaseModel, Field
from fastapi.responses import (
    FileResponse,
    JSONResponse,
    RedirectResponse,
    StreamingResponse
)
from datetime import date
from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from fastapi.staticfiles import StaticFiles
import base64
import hashlib
import hmac
import time
from urllib.parse import quote, urlencode

load_dotenv()

REMINDER_JOB_ID = "daily-dingtalk-reminder"


@asynccontextmanager
async def lifespan(app: FastAPI):
    reminder_hour = int(os.getenv("REMINDER_HOUR", "9"))
    reminder_minute = int(os.getenv("REMINDER_MINUTE", "0"))

    scheduler = AsyncIOScheduler(
        timezone="Asia/Shanghai"
    )

    scheduler.add_job(
        run_daily_dingtalk_reminders,
        trigger="cron",
        hour=reminder_hour,
        minute=reminder_minute,
        id=REMINDER_JOB_ID,
        replace_existing=True,
        coalesce=True,
        max_instances=1,
        misfire_grace_time=3600
    )

    scheduler.start()
    app.state.reminder_scheduler = scheduler

    try:
        yield
    finally:
        scheduler.shutdown(wait=False)


app = FastAPI(
    title="发票台账助手",
    version="0.1.0",
    lifespan=lifespan
)

STATIC_DIR = Path("static")

app.mount(
    "/static",
    StaticFiles(directory=STATIC_DIR),
    name="static"
)

DIFY_API_BASE_URL = os.getenv("DIFY_API_BASE_URL", "").rstrip("/")
DIFY_API_KEY = os.getenv("DIFY_API_KEY", "")
DIFY_ORDER_API_KEY = os.getenv("DIFY_ORDER_API_KEY", "")
DIFY_USER = os.getenv("DIFY_USER", "invoice-ledger-local")
DATA_DIR = Path("data")
DB_PATH = DATA_DIR / "invoice_ledger.db"
DINGTALK_WEBHOOK = os.getenv("DINGTALK_WEBHOOK", "").strip()
DINGTALK_SECRET = os.getenv("DINGTALK_SECRET", "").strip()
APP_USERNAME = os.getenv("APP_USERNAME", "").strip()
APP_PASSWORD = os.getenv("APP_PASSWORD", "")
APP_SESSION_SECRET = os.getenv("APP_SESSION_SECRET", "")
APP_SESSION_HOURS = max(
    1,
    int(os.getenv("APP_SESSION_HOURS", "12"))
)
APP_COOKIE_SECURE = (
    os.getenv("APP_COOKIE_SECURE", "false").lower()
    == "true"
)
SESSION_COOKIE_NAME = "ledger_session"


def get_db_connection():
    DATA_DIR.mkdir(exist_ok=True)

    connection = sqlite3.connect(DB_PATH)
    connection.row_factory = sqlite3.Row
    return connection


def init_database():
    with closing(get_db_connection()) as connection:
        connection.execute("""
            CREATE TABLE IF NOT EXISTS invoices (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                invoice_number TEXT NOT NULL UNIQUE,
                document_type TEXT NOT NULL,
                invoice_date TEXT,
                buyer_name TEXT,
                seller_name TEXT,
                amount_excl_tax REAL,
                tax_amount REAL,
                amount_incl_tax REAL,
                items_text TEXT,
                review_status TEXT,
                review_notes TEXT,
                prepayment_amount REAL NOT NULL DEFAULT 0,
                paid_amount REAL NOT NULL DEFAULT 0,
                due_date TEXT,
                payment_status TEXT NOT NULL DEFAULT '未付款',
                owner TEXT,
                source_filename TEXT,
                raw_dify_output TEXT,
                created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
                updated_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP
            )
        """)

        connection.execute("""
            CREATE TABLE IF NOT EXISTS invoice_order_links (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                invoice_id INTEGER NOT NULL,
                order_number TEXT NOT NULL,
                FOREIGN KEY (invoice_id) REFERENCES invoices(id),
                UNIQUE(invoice_id, order_number)
            )
        """)

        connection.execute("""
            CREATE TABLE IF NOT EXISTS notification_logs (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                invoice_id INTEGER NOT NULL,
                channel TEXT NOT NULL,
                event_key TEXT NOT NULL,
                reminder_status TEXT NOT NULL,
                due_date TEXT NOT NULL,
                sent_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
                FOREIGN KEY (invoice_id) REFERENCES invoices(id),
                UNIQUE(invoice_id, channel, event_key)
            )
        """)
        connection.execute("""
            CREATE TABLE IF NOT EXISTS order_notification_logs (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                order_id INTEGER NOT NULL,
                channel TEXT NOT NULL,
                event_key TEXT NOT NULL,
                reminder_status TEXT NOT NULL,
                due_date TEXT NOT NULL,
                sent_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
                FOREIGN KEY (order_id) REFERENCES orders(id),
                UNIQUE(order_id, channel, event_key)
            )
        """)
        connection.execute("""
            CREATE TABLE IF NOT EXISTS orders (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                order_number TEXT NOT NULL,
                order_type TEXT NOT NULL,
                contract_number TEXT,
                counterparty_name TEXT,
                order_date TEXT,
                delivery_date TEXT,
                order_amount REAL NOT NULL,
                settlement_method TEXT,
                payment_terms TEXT,
                prepayment_amount REAL NOT NULL DEFAULT 0,
                paid_amount REAL NOT NULL DEFAULT 0,
                payment_status TEXT NOT NULL,
                due_date TEXT,
                owner TEXT,
                review_status TEXT NOT NULL DEFAULT '待人工确认',
                review_notes TEXT,
                source_filename TEXT,
                raw_dify_output TEXT,
                created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
                updated_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
                UNIQUE(order_type, order_number)
            )
        """)

        connection.commit()


init_database()

class InvoiceCreate(BaseModel):
    document_type: str
    invoice_number: str
    invoice_date: str | None = None
    buyer_name: str | None = None
    seller_name: str | None = None
    amount_excl_tax: float | None = None
    tax_amount: float | None = None
    amount_incl_tax: float
    items_text: str | None = None
    order_numbers: list[str] = Field(default_factory=list)
    review_status: str = "待人工确认"
    review_notes: list[str] = Field(default_factory=list)

    # 由业务人员补充
    prepayment_amount: float = 0
    paid_amount: float = 0
    due_date: str | None = None
    owner: str | None = None

    # 留存原始识别数据，方便后续追溯
    source_filename: str | None = None
    raw_dify_output: dict[str, Any] = Field(default_factory=dict)


class OrderCreate(BaseModel):
    order_type: Literal["销售订单", "采购订单"]
    order_number: str
    contract_number: str | None = None
    counterparty_name: str | None = None
    order_date: str | None = None
    delivery_date: str | None = None
    order_amount: float = Field(gt=0)
    settlement_method: str | None = None
    payment_terms: str | None = None

    # 累计收付款包含预付款
    prepayment_amount: float = Field(default=0, ge=0)
    paid_amount: float = Field(default=0, ge=0)

    due_date: str | None = None
    owner: str | None = None
    review_status: str = "待人工确认"
    review_notes: list[str] = Field(default_factory=list)
    source_filename: str | None = None
    raw_dify_output: dict[str, Any] = Field(default_factory=dict)


class OrderUpdate(BaseModel):
    contract_number: str | None = None
    counterparty_name: str | None = None
    order_date: str | None = None
    delivery_date: str | None = None
    order_amount: float | None = Field(default=None, gt=0)
    settlement_method: str | None = None
    payment_terms: str | None = None

    prepayment_amount: float | None = Field(default=None, ge=0)
    paid_amount: float | None = Field(default=None, ge=0)
    due_date: str | None = None
    owner: str | None = None

    review_status: Literal[
        "待人工确认",
        "已确认",
        "需人工复核"
    ] | None = None

    review_notes: list[str] | None = None


class LoginRequest(BaseModel):
    username: str
    password: str


def encode_session_payload(data: dict) -> str:
    content = json.dumps(
        data,
        ensure_ascii=False,
        separators=(",", ":")
    ).encode("utf-8")

    return (
        base64.urlsafe_b64encode(content)
        .rstrip(b"=")
        .decode("ascii")
    )


def decode_session_payload(value: str) -> dict:
    padding = "=" * (-len(value) % 4)
    content = base64.urlsafe_b64decode(
        value + padding
    )
    return json.loads(content.decode("utf-8"))


def create_session_token(username: str) -> str:
    payload = encode_session_payload({
        "username": username,
        "expires_at": (
            int(time.time())
            + APP_SESSION_HOURS * 3600
        )
    })

    signature = hmac.new(
        APP_SESSION_SECRET.encode("utf-8"),
        payload.encode("ascii"),
        hashlib.sha256
    ).hexdigest()

    return f"{payload}.{signature}"


def verify_session_token(
    token: str | None
) -> str | None:
    if not token or not APP_SESSION_SECRET:
        return None

    try:
        payload, signature = token.split(
            ".",
            maxsplit=1
        )

        expected_signature = hmac.new(
            APP_SESSION_SECRET.encode("utf-8"),
            payload.encode("ascii"),
            hashlib.sha256
        ).hexdigest()

        if not secrets.compare_digest(
            signature,
            expected_signature
        ):
            return None

        session_data = decode_session_payload(
            payload
        )

        if int(
            session_data.get("expires_at", 0)
        ) < int(time.time()):
            return None

        username = session_data.get("username")

        if not isinstance(username, str):
            return None

        if not secrets.compare_digest(
            username.encode("utf-8"),
            APP_USERNAME.encode("utf-8")
        ):
            return None

        return username

    except (
        ValueError,
        TypeError,
        KeyError,
        UnicodeDecodeError,
        binascii.Error
    ):
        return None


def auth_is_configured() -> bool:
    return (
        bool(APP_USERNAME)
        and len(APP_PASSWORD) >= 8
        and len(APP_SESSION_SECRET) >= 32
    )


@app.middleware("http")
async def require_login(
    request: Request,
    call_next
):
    public_paths = {
        "/login",
        "/api/auth/login"
    }

    if request.url.path in public_paths:
        return await call_next(request)

    username = verify_session_token(
        request.cookies.get(SESSION_COOKIE_NAME)
    )

    if username:
        request.state.username = username
        return await call_next(request)

    if request.url.path.startswith("/api/"):
        return JSONResponse(
            status_code=401,
            content={
                "detail": "未登录或登录已过期。"
            }
        )

    next_url = request.url.path

    if request.url.query:
        next_url += f"?{request.url.query}"

    return RedirectResponse(
        url=f"/login?next={quote(next_url, safe='')}",
        status_code=303
    )


@app.get("/login", include_in_schema=False)
def login_page(request: Request):
    username = verify_session_token(
        request.cookies.get(SESSION_COOKIE_NAME)
    )

    if username:
        return RedirectResponse(
            url="/",
            status_code=303
        )

    return FileResponse(STATIC_DIR / "login.html")


@app.post("/api/auth/login")
def login(
    credentials: LoginRequest,
    response: Response
):
    if not auth_is_configured():
        raise HTTPException(
            status_code=500,
            detail=(
                "登录认证配置无效：密码至少8位，"
                "会话密钥至少32位。"
            )
        )

    username_correct = secrets.compare_digest(
        credentials.username.encode("utf-8"),
        APP_USERNAME.encode("utf-8")
    )

    password_correct = secrets.compare_digest(
        credentials.password.encode("utf-8"),
        APP_PASSWORD.encode("utf-8")
    )

    if not username_correct or not password_correct:
        raise HTTPException(
            status_code=401,
            detail="用户名或密码错误。"
        )

    max_age = APP_SESSION_HOURS * 3600

    response.set_cookie(
        key=SESSION_COOKIE_NAME,
        value=create_session_token(
            credentials.username
        ),
        max_age=max_age,
        httponly=True,
        secure=APP_COOKIE_SECURE,
        samesite="lax",
        path="/"
    )

    return {
        "status": "ok",
        "username": credentials.username,
        "expires_in": max_age
    }


@app.post("/api/auth/logout")
def logout(response: Response):
    response.delete_cookie(
        key=SESSION_COOKIE_NAME,
        path="/"
    )

    return {
        "status": "ok",
        "message": "已退出登录。"
    }


@app.get("/", include_in_schema=False)
def home():
    return FileResponse(STATIC_DIR / "index.html")

@app.get("/api/health")
def health_check():
    return {
        "status": "ok",
        "dify_base_url": DIFY_API_BASE_URL,
        "dify_invoice_api_key_configured": bool(DIFY_API_KEY),
        "dify_order_api_key_configured": bool(DIFY_ORDER_API_KEY)
    }


@app.post("/api/dify/extract")
async def extract_invoice(
    document_type: str = Form("销售发票"),
    invoice_file: UploadFile = File(...)
):
    if not DIFY_API_BASE_URL or not DIFY_API_KEY:
        raise HTTPException(
            status_code=500,
            detail="Dify 配置缺失，请检查 .env 文件。"
        )

    content_type = invoice_file.content_type or "application/octet-stream"

    if content_type.startswith("image/"):
        dify_file_type = "image"
    elif content_type == "application/pdf":
        dify_file_type = "document"
    else:
        raise HTTPException(
            status_code=400,
            detail="当前只支持图片和 PDF 文件。"
        )

    file_content = await invoice_file.read()

    headers = {
        "Authorization": f"Bearer {DIFY_API_KEY}"
    }

    try:
        async with httpx.AsyncClient(timeout=90) as client:
            # 第一步：将文件上传到 Dify，获取 upload_file_id
            upload_response = await client.post(
                f"{DIFY_API_BASE_URL}/files/upload",
                headers=headers,
                data={"user": DIFY_USER},
                files={
                    "file": (
                        invoice_file.filename,
                        file_content,
                        content_type
                    )
                }
            )
            upload_response.raise_for_status()
            upload_file_id = upload_response.json()["id"]

            # 第二步：调用已发布的 Workflow
            workflow_payload = {
                "inputs": {
                    "document_type": document_type,
                    "invoice_file": [
                        {
                            "transfer_method": "local_file",
                            "upload_file_id": upload_file_id,
                            "type": dify_file_type
                        }
                    ]
                },
                "response_mode": "blocking",
                "user": DIFY_USER
            }

            workflow_response = await client.post(
                f"{DIFY_API_BASE_URL}/workflows/run",
                headers={
                    **headers,
                    "Content-Type": "application/json"
                },
                json=workflow_payload
            )
            workflow_response.raise_for_status()
            workflow_result = workflow_response.json()

    except httpx.HTTPStatusError as error:
        try:
            detail = error.response.json()
        except ValueError:
            detail = error.response.text

        raise HTTPException(
            status_code=error.response.status_code,
            detail=detail
        ) from error

    return {
        "workflow_run_id": workflow_result.get("workflow_run_id"),
        "status": workflow_result.get("data", {}).get("status"),
        "outputs": workflow_result.get("data", {}).get("outputs", {})
    }


@app.post("/api/dify/orders/extract")
async def extract_order(
    order_type: str = Form(...),
    order_file: UploadFile = File(...)
):
    if order_type not in {"销售订单", "采购订单"}:
        raise HTTPException(
            status_code=422,
            detail="order_type 必须是销售订单或采购订单。"
        )

    if not DIFY_API_BASE_URL or not DIFY_ORDER_API_KEY:
        raise HTTPException(
            status_code=500,
            detail="订单 Dify 配置缺失，请检查 .env 文件。"
        )

    content_type = (
        order_file.content_type
        or "application/octet-stream"
    )

    # 当前订单 Workflow 暂时只接收图片
    if not content_type.startswith("image/"):
        raise HTTPException(
            status_code=400,
            detail="当前订单智能识别只支持 PNG、JPG 等图片文件。"
        )

    file_content = await order_file.read()

    if not file_content:
        raise HTTPException(
            status_code=400,
            detail="上传的订单文件为空。"
        )

    headers = {
        "Authorization": f"Bearer {DIFY_ORDER_API_KEY}"
    }

    try:
        async with httpx.AsyncClient(timeout=90) as client:
            # 第一步：将订单图片上传到订单 Workflow
            upload_response = await client.post(
                f"{DIFY_API_BASE_URL}/files/upload",
                headers=headers,
                data={"user": DIFY_USER},
                files={
                    "file": (
                        order_file.filename,
                        file_content,
                        content_type
                    )
                }
            )
            upload_response.raise_for_status()

            upload_file_id = upload_response.json()["id"]

            # 第二步：执行订单识别 Workflow
            workflow_payload = {
                "inputs": {
                    "order_type": order_type,
                    "order_file": [
                        {
                            "transfer_method": "local_file",
                            "upload_file_id": upload_file_id,
                            "type": "image"
                        }
                    ]
                },
                "response_mode": "blocking",
                "user": DIFY_USER
            }

            workflow_response = await client.post(
                f"{DIFY_API_BASE_URL}/workflows/run",
                headers={
                    **headers,
                    "Content-Type": "application/json"
                },
                json=workflow_payload
            )
            workflow_response.raise_for_status()
            workflow_result = workflow_response.json()

    except httpx.HTTPStatusError as error:
        try:
            detail = error.response.json()
        except ValueError:
            detail = error.response.text

        raise HTTPException(
            status_code=error.response.status_code,
            detail=detail
        ) from error

    except httpx.RequestError as error:
        raise HTTPException(
            status_code=502,
            detail=(
                "无法连接到 Dify，请确认 SSH 隧道仍在运行："
                f"{str(error)}"
            )
        ) from error

    return {
        "workflow_run_id": workflow_result.get(
            "workflow_run_id"
        ),
        "status": workflow_result.get(
            "data", {}
        ).get("status"),
        "outputs": workflow_result.get(
            "data", {}
        ).get("outputs", {})
    }

@app.post("/api/invoices", status_code=201)
def create_invoice(invoice: InvoiceCreate):
    if invoice.paid_amount >= invoice.amount_incl_tax:
        payment_status = "已收款" if invoice.document_type == "销售发票" else "已付款"
    elif invoice.paid_amount > 0:
        payment_status = "部分收款" if invoice.document_type == "销售发票" else "部分付款"
    else:
        payment_status = "未收款" if invoice.document_type == "销售发票" else "未付款"

    try:
        with closing(get_db_connection()) as connection:
            cursor = connection.execute(
                """
                INSERT INTO invoices (
                    invoice_number,
                    document_type,
                    invoice_date,
                    buyer_name,
                    seller_name,
                    amount_excl_tax,
                    tax_amount,
                    amount_incl_tax,
                    items_text,
                    review_status,
                    review_notes,
                    prepayment_amount,
                    paid_amount,
                    due_date,
                    payment_status,
                    owner,
                    source_filename,
                    raw_dify_output
                )
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                (
                    invoice.invoice_number,
                    invoice.document_type,
                    invoice.invoice_date,
                    invoice.buyer_name,
                    invoice.seller_name,
                    invoice.amount_excl_tax,
                    invoice.tax_amount,
                    invoice.amount_incl_tax,
                    invoice.items_text,
                    invoice.review_status,
                    json.dumps(invoice.review_notes, ensure_ascii=False),
                    invoice.prepayment_amount,
                    invoice.paid_amount,
                    invoice.due_date,
                    payment_status,
                    invoice.owner,
                    invoice.source_filename,
                    json.dumps(invoice.raw_dify_output, ensure_ascii=False)
                )
            )

            invoice_id = cursor.lastrowid

            for order_number in set(invoice.order_numbers):
                if order_number.strip():
                    connection.execute(
                        """
                        INSERT INTO invoice_order_links (invoice_id, order_number)
                        VALUES (?, ?)
                        """,
                        (invoice_id, order_number.strip())
                    )

            connection.commit()

    except sqlite3.IntegrityError as error:
        raise HTTPException(
            status_code=409,
            detail="该发票号码已存在，不能重复保存。"
        ) from error

    return {
        "id": invoice_id,
        "invoice_number": invoice.invoice_number,
        "payment_status": payment_status,
        "message": "发票已保存到台账。"
    }

@app.get("/api/invoices")
def list_invoices():
    with closing(get_db_connection()) as connection:
        rows = connection.execute(
            """
            SELECT
                invoices.*,
                GROUP_CONCAT(invoice_order_links.order_number) AS order_numbers_text
            FROM invoices
            LEFT JOIN invoice_order_links
                ON invoices.id = invoice_order_links.invoice_id
            GROUP BY invoices.id
            ORDER BY invoices.created_at DESC
            """
        ).fetchall()

    items = []

    for row in rows:
        invoice = dict(row)

        order_numbers_text = invoice.pop("order_numbers_text", None)
        invoice["order_numbers"] = (
            order_numbers_text.split(",")
            if order_numbers_text
            else []
        )

        invoice["review_notes"] = json.loads(
            invoice["review_notes"] or "[]"
        )
        invoice["raw_dify_output"] = json.loads(
            invoice["raw_dify_output"] or "{}"
        )

        items.append(invoice)

    return {
        "total": len(items),
        "items": items
    }

class InvoiceUpdate(BaseModel):
    prepayment_amount: float | None = Field(default=None, ge=0)
    paid_amount: float | None = Field(default=None, ge=0)
    due_date: str | None = None
    owner: str | None = None
    review_status: str | None = None
    review_notes: list[str] | None = None


@app.patch("/api/invoices/{invoice_id}")
def update_invoice(invoice_id: int, invoice_update: InvoiceUpdate):
    update_data = invoice_update.model_dump(exclude_unset=True)

    if not update_data:
        raise HTTPException(
            status_code=400,
            detail="请至少提交一个需要更新的字段。"
        )

    with closing(get_db_connection()) as connection:
        existing = connection.execute(
            "SELECT * FROM invoices WHERE id = ?",
            (invoice_id,)
        ).fetchone()

        if existing is None:
            raise HTTPException(
                status_code=404,
                detail="未找到该发票台账。"
            )

        paid_amount = update_data.get(
            "paid_amount",
            existing["paid_amount"]
        )
        prepayment_amount = update_data.get(
            "prepayment_amount",
            existing["prepayment_amount"]
        )
        amount_incl_tax = existing["amount_incl_tax"]

        if paid_amount > amount_incl_tax:
            raise HTTPException(
                status_code=422,
                detail="累计收付款金额不能大于发票价税合计。"
            )

        if prepayment_amount > paid_amount:
            raise HTTPException(
                status_code=422,
                detail="预付款金额不能大于累计收付款金额。"
            )

        if paid_amount >= amount_incl_tax:
            payment_status = (
                "已收款"
                if existing["document_type"] == "销售发票"
                else "已付款"
            )
        elif paid_amount > 0:
            payment_status = (
                "部分收款"
                if existing["document_type"] == "销售发票"
                else "部分付款"
            )
        else:
            payment_status = (
                "未收款"
                if existing["document_type"] == "销售发票"
                else "未付款"
            )

        update_data["payment_status"] = payment_status

        if "review_notes" in update_data:
            update_data["review_notes"] = json.dumps(
                update_data["review_notes"],
                ensure_ascii=False
            )

        assignments = ", ".join(
            f"{field} = ?"
            for field in update_data
        )
        values = list(update_data.values())
        values.append(invoice_id)

        connection.execute(
            f"""
            UPDATE invoices
            SET {assignments},
                updated_at = CURRENT_TIMESTAMP
            WHERE id = ?
            """,
            values
        )
        connection.commit()

    return {
        "id": invoice_id,
        "payment_status": payment_status,
        "message": "台账已更新。"
    }

def get_due_reminders(days_before: int):
    today = date.today()
    reminders = []

    with closing(get_db_connection()) as connection:
        rows = connection.execute(
            """
            SELECT
                invoices.*,
                GROUP_CONCAT(
                    invoice_order_links.order_number
                ) AS order_numbers_text
            FROM invoices
            LEFT JOIN invoice_order_links
                ON invoice_order_links.invoice_id =
                   invoices.id
            GROUP BY invoices.id
            """
        ).fetchall()

    for row in rows:
        invoice = dict(row)

        order_numbers_text = invoice.pop(
            "order_numbers_text",
            None
        )

        order_numbers = (
            order_numbers_text.split(",")
            if order_numbers_text
            else []
        )

        if not invoice["due_date"]:
            continue

        try:
            due_date = date.fromisoformat(invoice["due_date"])
        except ValueError:
            continue

        amount_due = round(
            max(
                0,
                invoice["amount_incl_tax"] - invoice["paid_amount"]
            ),
            2
        )

        # 已结清的单据不提醒
        if amount_due <= 0.01:
            continue

        days_to_due = (due_date - today).days

        if days_to_due < 0:
            reminder_status = "已逾期"
        elif days_to_due == 0:
            reminder_status = "今日到期"
        elif days_to_due <= days_before:
            reminder_status = "即将到期"
        else:
            continue

        counterparty = (
            invoice["buyer_name"]
            if invoice["document_type"] == "销售发票"
            else invoice["seller_name"]
        )

        reminders.append({
            "source_type": "invoice",
            "invoice_id": invoice["id"],
            "invoice_number": invoice["invoice_number"],
            "document_type": invoice["document_type"],
            "order_numbers": order_numbers,
            "counterparty": counterparty,
            "due_date": invoice["due_date"],
            "days_to_due": days_to_due,
            "amount_due": amount_due,
            "payment_status": invoice["payment_status"],
            "owner": invoice["owner"],
            "reminder_status": reminder_status
        })

    return sorted(
        reminders,
        key=lambda item: item["days_to_due"]
    )

def get_order_due_reminders(days_before: int):
    today = date.today()
    reminders = []

    with closing(get_db_connection()) as connection:
        rows = connection.execute(
            "SELECT * FROM orders"
        ).fetchall()

    for row in rows:
        order = dict(row)

        # 没填写到期日的订单暂不提醒
        if not order["due_date"]:
            continue

        try:
            due_date_value = date.fromisoformat(
                order["due_date"]
            )
        except ValueError:
            continue

        amount_due = round(
            max(
                0,
                order["order_amount"] - order["paid_amount"]
            ),
            2
        )

        # 已经完全收款或付款的订单不提醒
        if amount_due <= 0.01:
            continue

        days_to_due = (
            due_date_value - today
        ).days

        if days_to_due < 0:
            reminder_status = "已逾期"
        elif days_to_due == 0:
            reminder_status = "今日到期"
        elif days_to_due <= days_before:
            reminder_status = "即将到期"
        else:
            continue

        reminders.append({
            "source_type": "order",
            "order_id": order["id"],
            "order_number": order["order_number"],
            "order_type": order["order_type"],
            "counterparty": order["counterparty_name"],
            "due_date": order["due_date"],
            "days_to_due": days_to_due,
            "amount_due": amount_due,
            "payment_status": order["payment_status"],
            "owner": order["owner"],
            "reminder_status": reminder_status
        })

    return sorted(
        reminders,
        key=lambda item: item["days_to_due"]
    )

def get_combined_due_reminders(
    days_before: int
) -> list[dict]:
    invoice_reminders = get_due_reminders(
        days_before
    )

    order_reminders = get_order_due_reminders(
        days_before
    )

    # 记录已由发票承担提醒的订单
    covered_orders = set()

    for reminder in invoice_reminders:
        if reminder["document_type"] == "销售发票":
            matching_order_type = "销售订单"
        elif reminder["document_type"] == "采购发票":
            matching_order_type = "采购订单"
        else:
            continue

        for order_number in reminder.get(
            "order_numbers",
            []
        ):
            covered_orders.add(
                (
                    matching_order_type,
                    order_number
                )
            )

    filtered_order_reminders = []

    for reminder in order_reminders:
        order_key = (
            reminder["order_type"],
            reminder["order_number"]
        )

        if order_key in covered_orders:
            continue

        filtered_order_reminders.append(
            reminder
        )

    combined = (
        invoice_reminders
        + filtered_order_reminders
    )

    return sorted(
        combined,
        key=lambda item: item["days_to_due"]
    )

@app.get("/api/reminders/all/due")
def list_all_due_reminders(
    days_before: int = 7
):
    if days_before < 0 or days_before > 90:
        raise HTTPException(
            status_code=422,
            detail="days_before 必须在 0 到 90 之间。"
        )

    invoice_reminders = get_due_reminders(
        days_before
    )

    order_reminders = get_order_due_reminders(
        days_before
    )

    combined = get_combined_due_reminders(
        days_before
    )

    return {
        "today": date.today().isoformat(),
        "days_before": days_before,
        "invoice_total_before_dedup": len(
            invoice_reminders
        ),
        "order_total_before_dedup": len(
            order_reminders
        ),
        "suppressed_duplicate_total": (
            len(invoice_reminders)
            + len(order_reminders)
            - len(combined)
        ),
        "total": len(combined),
        "items": combined
    }
@app.get("/api/reminders/due")
def list_due_reminders(days_before: int = 7):
    if days_before < 0 or days_before > 90:
        raise HTTPException(
            status_code=422,
            detail="days_before 必须在 0 到 90 之间。"
        )

    reminders = get_due_reminders(days_before)

    return {
        "today": date.today().isoformat(),
        "days_before": days_before,
        "total": len(reminders),
        "items": reminders
    }

@app.get("/api/reminders/orders/due")
def list_order_due_reminders(
    days_before: int = 7
):
    if days_before < 0 or days_before > 90:
        raise HTTPException(
            status_code=422,
            detail="days_before 必须在 0 到 90 之间。"
        )

    reminders = get_order_due_reminders(
        days_before
    )

    return {
        "today": date.today().isoformat(),
        "days_before": days_before,
        "total": len(reminders),
        "items": reminders
    }

def build_dingtalk_signed_url() -> str:
    if not DINGTALK_WEBHOOK or not DINGTALK_SECRET:
        raise HTTPException(
            status_code=500,
            detail="钉钉配置缺失，请检查 .env 文件。"
        )

    timestamp = str(int(time.time() * 1000))
    string_to_sign = f"{timestamp}\n{DINGTALK_SECRET}"

    signature = hmac.new(
        DINGTALK_SECRET.encode("utf-8"),
        string_to_sign.encode("utf-8"),
        digestmod=hashlib.sha256
    ).digest()

    sign = base64.b64encode(signature).decode("utf-8")

    separator = "&" if "?" in DINGTALK_WEBHOOK else "?"

    return (
        f"{DINGTALK_WEBHOOK}"
        f"{separator}"
        f"{urlencode({'timestamp': timestamp, 'sign': sign})}"
    )


async def send_dingtalk_text(content: str) -> dict:
    signed_url = build_dingtalk_signed_url()

    payload = {
        "msgtype": "text",
        "text": {
            "content": content
        }
    }

    try:
        async with httpx.AsyncClient(timeout=15) as client:
            response = await client.post(
                signed_url,
                json=payload
            )
            response.raise_for_status()

    except httpx.HTTPError as error:
        raise HTTPException(
            status_code=502,
            detail=f"调用钉钉接口失败：{str(error)}"
        ) from error

    try:
        result = response.json()
    except ValueError as error:
        raise HTTPException(
            status_code=502,
            detail="钉钉返回了无法解析的响应。"
        ) from error

    if result.get("errcode") != 0:
        raise HTTPException(
            status_code=502,
            detail={
                "message": "钉钉发送失败",
                "dingtalk_response": result
            }
        )

    return result


@app.post("/api/notifications/dingtalk/test")
async def test_dingtalk_notification():
    content = (
        "【发票台账助手】钉钉通知测试\n"
        f"测试日期：{date.today().isoformat()}\n"
        "状态：Webhook 与加签配置正常。"
    )

    result = await send_dingtalk_text(content)

    return {
        "status": "sent",
        "message": "钉钉测试消息已发送。",
        "dingtalk_response": result
    }

def build_due_reminder_message(
    reminders: list[dict]
) -> str:
    lines = [
        "【供应链账期提醒】",
        f"统计日期：{date.today().isoformat()}",
        f"待跟进单据：{len(reminders)} 笔",
        ""
    ]

    for index, item in enumerate(
        reminders,
        start=1
    ):
        days_to_due = item["days_to_due"]

        if days_to_due < 0:
            time_description = (
                f"已逾期 {abs(days_to_due)} 天"
            )
        elif days_to_due == 0:
            time_description = "今日到期"
        else:
            time_description = (
                f"剩余 {days_to_due} 天"
            )

        if item["source_type"] == "invoice":
            document_type = item["document_type"]
            number_label = "发票号码"
            document_number = item["invoice_number"]

            order_numbers = item.get(
                "order_numbers",
                []
            )
        else:
            document_type = item["order_type"]
            number_label = "订单号码"
            document_number = item["order_number"]
            order_numbers = []

        is_receivable = document_type in {
            "销售发票",
            "销售订单"
        }

        amount_label = (
            "待收款"
            if is_receivable
            else "待付款"
        )

        lines.extend([
            (
                f"{index}. "
                f"{item['reminder_status']}｜"
                f"{document_type}"
            ),
            (
                "往来单位："
                f"{item['counterparty'] or '待确认'}"
            ),
            (
                f"{number_label}："
                f"{document_number}"
            ),
            (
                "到期日期："
                f"{item['due_date']}"
                f"（{time_description}）"
            ),
            (
                f"{amount_label}："
                f"¥{item['amount_due']:,.2f}"
            ),
            (
                "负责人："
                f"{item['owner'] or '待分配'}"
            )
        ])

        if order_numbers:
            lines.append(
                "关联订单："
                + "、".join(order_numbers)
            )

        lines.append("")

    lines.append("请相关负责人及时跟进并更新台账。")

    return "\n".join(lines)


@app.post("/api/notifications/dingtalk/reminders")
async def send_due_reminder_notification(days_before: int = 7):
    if days_before < 0 or days_before > 90:
        raise HTTPException(
            status_code=422,
            detail="days_before 必须在 0 到 90 之间。"
        )

    reminders = get_combined_due_reminders(
        days_before
    )

    if not reminders:
        return {
            "status": "skipped",
            "sent_count": 0,
            "message": "当前范围内没有需要提醒的未结清单据。"
        }

    content = build_due_reminder_message(reminders)
    dingtalk_result = await send_dingtalk_text(content)

    return {
        "status": "sent",
        "sent_count": len(reminders),
        "message": "账期提醒已发送到钉钉群。",
        "dingtalk_response": dingtalk_result
    }

@app.post("/api/orders", status_code=201)
def create_order(order: OrderCreate):
    order_number = order.order_number.strip()

    if not order_number:
        raise HTTPException(
            status_code=422,
            detail="订单号不能为空。"
        )

    if order.paid_amount > order.order_amount:
        raise HTTPException(
            status_code=422,
            detail="累计收付款金额不能大于订单金额。"
        )

    if order.prepayment_amount > order.paid_amount:
        raise HTTPException(
            status_code=422,
            detail="预付款不能大于累计收付款金额。"
        )

    if order.paid_amount >= order.order_amount:
        payment_status = (
            "已收款"
            if order.order_type == "销售订单"
            else "已付款"
        )
    elif order.paid_amount > 0:
        payment_status = (
            "部分收款"
            if order.order_type == "销售订单"
            else "部分付款"
        )
    else:
        payment_status = (
            "未收款"
            if order.order_type == "销售订单"
            else "未付款"
        )

    try:
        with closing(get_db_connection()) as connection:
            cursor = connection.execute(
                """
                INSERT INTO orders (
                    order_number,
                    order_type,
                    contract_number,
                    counterparty_name,
                    order_date,
                    delivery_date,
                    order_amount,
                    settlement_method,
                    payment_terms,
                    prepayment_amount,
                    paid_amount,
                    payment_status,
                    due_date,
                    owner,
                    review_status,
                    review_notes,
                    source_filename,
                    raw_dify_output
                )
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                (
                    order_number,
                    order.order_type,
                    order.contract_number,
                    order.counterparty_name,
                    order.order_date,
                    order.delivery_date,
                    order.order_amount,
                    order.settlement_method,
                    order.payment_terms,
                    order.prepayment_amount,
                    order.paid_amount,
                    payment_status,
                    order.due_date,
                    order.owner,
                    order.review_status,
                    json.dumps(order.review_notes, ensure_ascii=False),
                    order.source_filename,
                    json.dumps(order.raw_dify_output, ensure_ascii=False)
                )
            )

            order_id = cursor.lastrowid
            connection.commit()

    except sqlite3.IntegrityError as error:
        raise HTTPException(
            status_code=409,
            detail="相同类型的订单号已经存在，不能重复保存。"
        ) from error

    return {
        "id": order_id,
        "order_number": order_number,
        "payment_status": payment_status,
        "message": "订单已保存到台账。"
    }

@app.get("/api/orders")
def list_orders():
    with closing(get_db_connection()) as connection:
        rows = connection.execute(
            """
            SELECT
                orders.*,
                GROUP_CONCAT(
                    DISTINCT invoices.invoice_number
                ) AS invoice_numbers_text
            FROM orders
            LEFT JOIN invoice_order_links
                ON invoice_order_links.order_number =
                   orders.order_number
            LEFT JOIN invoices
                ON invoices.id =
                   invoice_order_links.invoice_id
               AND (
                    (
                        orders.order_type = '销售订单'
                        AND invoices.document_type = '销售发票'
                    )
                    OR
                    (
                        orders.order_type = '采购订单'
                        AND invoices.document_type = '采购发票'
                    )
               )
            GROUP BY orders.id
            ORDER BY orders.created_at DESC
            """
        ).fetchall()

    items = []

    for row in rows:
        order = dict(row)

        invoice_numbers_text = order.pop(
            "invoice_numbers_text",
            None
        )

        order["invoice_numbers"] = (
            invoice_numbers_text.split(",")
            if invoice_numbers_text
            else []
        )

        order["invoice_status"] = (
            "已关联发票"
            if order["invoice_numbers"]
            else "未开票"
        )

        order["review_notes"] = json.loads(
            order["review_notes"] or "[]"
        )
        order["raw_dify_output"] = json.loads(
            order["raw_dify_output"] or "{}"
        )

        items.append(order)

    return {
        "total": len(items),
        "items": items
    }


@app.patch("/api/orders/{order_id}")
def update_order(
    order_id: int,
    order_update: OrderUpdate
):
    update_data = order_update.model_dump(
        exclude_unset=True
    )

    if not update_data:
        raise HTTPException(
            status_code=400,
            detail="请至少提交一个需要更新的字段。"
        )

    numeric_fields = {
        "order_amount",
        "prepayment_amount",
        "paid_amount"
    }

    for field in numeric_fields:
        if field in update_data and update_data[field] is None:
            raise HTTPException(
                status_code=422,
                detail=f"{field} 不能设置为空。"
            )

    with closing(get_db_connection()) as connection:
        existing = connection.execute(
            "SELECT * FROM orders WHERE id = ?",
            (order_id,)
        ).fetchone()

        if existing is None:
            raise HTTPException(
                status_code=404,
                detail="未找到该订单台账。"
            )

        order_amount = update_data.get(
            "order_amount",
            existing["order_amount"]
        )

        paid_amount = update_data.get(
            "paid_amount",
            existing["paid_amount"]
        )

        prepayment_amount = update_data.get(
            "prepayment_amount",
            existing["prepayment_amount"]
        )

        if paid_amount > order_amount:
            raise HTTPException(
                status_code=422,
                detail="累计收付款金额不能大于订单金额。"
            )

        if prepayment_amount > paid_amount:
            raise HTTPException(
                status_code=422,
                detail="预付款金额不能大于累计收付款金额。"
            )

        if paid_amount >= order_amount:
            payment_status = (
                "已收款"
                if existing["order_type"] == "销售订单"
                else "已付款"
            )
        elif paid_amount > 0:
            payment_status = (
                "部分收款"
                if existing["order_type"] == "销售订单"
                else "部分付款"
            )
        else:
            payment_status = (
                "未收款"
                if existing["order_type"] == "销售订单"
                else "未付款"
            )

        update_data["payment_status"] = payment_status

        if "review_notes" in update_data:
            update_data["review_notes"] = json.dumps(
                update_data["review_notes"] or [],
                ensure_ascii=False
            )

        assignments = ", ".join(
            f"{field} = ?"
            for field in update_data
        )

        values = list(update_data.values())
        values.append(order_id)

        connection.execute(
            f"""
            UPDATE orders
            SET {assignments},
                updated_at = CURRENT_TIMESTAMP
            WHERE id = ?
            """,
            values
        )

        connection.commit()

    return {
        "id": order_id,
        "payment_status": payment_status,
        "message": "订单台账已更新。"
    }

def get_notification_event_key(reminder: dict) -> str | None:
    days_to_due = reminder["days_to_due"]
    due_date = reminder["due_date"]

    if 4 <= days_to_due <= 7:
        return f"{due_date}:D7"

    if 1 <= days_to_due <= 3:
        return f"{due_date}:D3"

    if days_to_due == 0:
        return f"{due_date}:D0"

    if days_to_due < 0:
        return (
            f"{due_date}:OVERDUE:"
            f"{date.today().isoformat()}"
        )

    return None


def get_unsent_dingtalk_reminders(
    reminders: list[dict]
) -> list[dict]:
    unsent = []

    with closing(get_db_connection()) as connection:
        for reminder in reminders:
            event_key = get_notification_event_key(
                reminder
            )

            if event_key is None:
                continue

            source_type = reminder.get(
                "source_type"
            )

            if source_type == "invoice":
                existing = connection.execute(
                    """
                    SELECT id
                    FROM notification_logs
                    WHERE invoice_id = ?
                      AND channel = ?
                      AND event_key = ?
                    """,
                    (
                        reminder["invoice_id"],
                        "dingtalk",
                        event_key
                    )
                ).fetchone()

            elif source_type == "order":
                existing = connection.execute(
                    """
                    SELECT id
                    FROM order_notification_logs
                    WHERE order_id = ?
                      AND channel = ?
                      AND event_key = ?
                    """,
                    (
                        reminder["order_id"],
                        "dingtalk",
                        event_key
                    )
                ).fetchone()

            else:
                continue

            if existing is None:
                reminder_with_event = dict(reminder)
                reminder_with_event["event_key"] = (
                    event_key
                )
                unsent.append(
                    reminder_with_event
                )

    return unsent

def record_dingtalk_notifications(
    reminders: list[dict]
) -> None:
    with closing(get_db_connection()) as connection:
        for reminder in reminders:
            if reminder["source_type"] == "invoice":
                connection.execute(
                    """
                    INSERT OR IGNORE INTO notification_logs (
                        invoice_id,
                        channel,
                        event_key,
                        reminder_status,
                        due_date
                    )
                    VALUES (?, ?, ?, ?, ?)
                    """,
                    (
                        reminder["invoice_id"],
                        "dingtalk",
                        reminder["event_key"],
                        reminder["reminder_status"],
                        reminder["due_date"]
                    )
                )

            elif reminder["source_type"] == "order":
                connection.execute(
                    """
                    INSERT OR IGNORE INTO order_notification_logs (
                        order_id,
                        channel,
                        event_key,
                        reminder_status,
                        due_date
                    )
                    VALUES (?, ?, ?, ?, ?)
                    """,
                    (
                        reminder["order_id"],
                        "dingtalk",
                        reminder["event_key"],
                        reminder["reminder_status"],
                        reminder["due_date"]
                    )
                )

        connection.commit()


@app.post("/api/notifications/dingtalk/run-daily")
async def run_daily_dingtalk_reminders():
    reminders = get_combined_due_reminders(
        days_before=7
    )
    unsent_reminders = get_unsent_dingtalk_reminders(reminders)

    if not unsent_reminders:
        return {
            "status": "skipped",
            "sent_count": 0,
            "message": "今天没有新的账期提醒需要发送。"
        }

    content = build_due_reminder_message(unsent_reminders)
    dingtalk_result = await send_dingtalk_text(content)

    record_dingtalk_notifications(unsent_reminders)

    return {
        "status": "sent",
        "sent_count": len(unsent_reminders),
        "message": "今日账期提醒已发送。",
        "dingtalk_response": dingtalk_result
    }
@app.get("/api/scheduler/status")
def get_scheduler_status():
    scheduler = getattr(
        app.state,
        "reminder_scheduler",
        None
    )

    if scheduler is None:
        return {
            "running": False,
            "message": "提醒调度器尚未启动。"
        }

    job = scheduler.get_job(REMINDER_JOB_ID)

    return {
        "running": scheduler.running,
        "job_id": REMINDER_JOB_ID,
        "next_run_time": (
            job.next_run_time.isoformat()
            if job and job.next_run_time
            else None
        )
    }


@app.get("/api/export/ledger.xlsx")
def export_ledger_excel():
    order_result = list_orders()
    invoice_result = list_invoices()

    orders = order_result["items"]
    invoices = invoice_result["items"]

    workbook = Workbook()

    # 工作表一：订单发票总表
    order_sheet = workbook.active
    order_sheet.title = "订单发票总表"

    order_sheet.append([
        "订单类型",
        "订单号",
        "合同号",
        "往来单位",
        "订单日期",
        "交付日期",
        "订单金额",
        "预付款金额",
        "累计收付款",
        "剩余金额",
        "到期日",
        "收付款状态",
        "开票状态",
        "关联发票",
        "结算方式",
        "付款账期",
        "负责人",
        "复核状态"
    ])

    for order in orders:
        order_sheet.append([
            order["order_type"],
            order["order_number"],
            order["contract_number"],
            order["counterparty_name"],
            parse_excel_date(order["order_date"]),
            parse_excel_date(order["delivery_date"]),
            float(order["order_amount"] or 0),
            float(order["prepayment_amount"] or 0),
            float(order["paid_amount"] or 0),
            None,
            parse_excel_date(order["due_date"]),
            order["payment_status"],
            order["invoice_status"],
            "、".join(order["invoice_numbers"] or []),
            order["settlement_method"],
            order["payment_terms"],
            order["owner"],
            order["review_status"]
        ])

        row_number = order_sheet.max_row

        # 剩余金额 = 订单金额 - 累计收付款
        order_sheet.cell(
            row=row_number,
            column=10,
            value=f"=MAX(0,G{row_number}-I{row_number})"
        )

    style_export_sheet(
        order_sheet,
        [
            12, 20, 20, 28, 13, 13,
            15, 15, 15, 15, 13, 14,
            14, 30, 20, 32, 14, 14
        ]
    )

    for row_number in range(
        2,
        order_sheet.max_row + 1
    ):
        for column in (7, 8, 9, 10):
            order_sheet.cell(
                row=row_number,
                column=column
            ).number_format = "#,##0.00"

        for column in (5, 6, 11):
            order_sheet.cell(
                row=row_number,
                column=column
            ).number_format = "yyyy-mm-dd"

    # 工作表二：发票明细
    invoice_sheet = workbook.create_sheet(
        "发票明细"
    )

    invoice_sheet.append([
        "发票类型",
        "发票号码",
        "开票日期",
        "购买方",
        "销售方",
        "未税金额",
        "税额",
        "价税合计",
        "预付款金额",
        "累计收付款",
        "剩余金额",
        "到期日",
        "收付款状态",
        "关联订单",
        "负责人",
        "复核状态",
        "来源文件"
    ])

    for invoice in invoices:
        invoice_sheet.append([
            invoice["document_type"],
            invoice["invoice_number"],
            parse_excel_date(invoice["invoice_date"]),
            invoice["buyer_name"],
            invoice["seller_name"],
            float(invoice["amount_excl_tax"] or 0),
            float(invoice["tax_amount"] or 0),
            float(invoice["amount_incl_tax"] or 0),
            float(invoice["prepayment_amount"] or 0),
            float(invoice["paid_amount"] or 0),
            None,
            parse_excel_date(invoice["due_date"]),
            invoice["payment_status"],
            "、".join(invoice["order_numbers"] or []),
            invoice["owner"],
            invoice["review_status"],
            invoice["source_filename"]
        ])

        row_number = invoice_sheet.max_row

        # 剩余金额 = 价税合计 - 累计收付款
        invoice_sheet.cell(
            row=row_number,
            column=11,
            value=f"=MAX(0,H{row_number}-J{row_number})"
        )

    style_export_sheet(
        invoice_sheet,
        [
            12, 24, 13, 30, 30, 15,
            15, 15, 15, 15, 15, 13,
            14, 32, 14, 14, 30
        ]
    )

    for row_number in range(
        2,
        invoice_sheet.max_row + 1
    ):
        for column in (6, 7, 8, 9, 10, 11):
            invoice_sheet.cell(
                row=row_number,
                column=column
            ).number_format = "#,##0.00"

        for column in (3, 12):
            invoice_sheet.cell(
                row=row_number,
                column=column
            ).number_format = "yyyy-mm-dd"

    output = BytesIO()
    workbook.save(output)
    output.seek(0)

    filename = (
        "supply_chain_ledger_"
        f"{date.today().strftime('%Y%m%d')}.xlsx"
    )

    return StreamingResponse(
        output,
        media_type=(
            "application/vnd.openxmlformats-"
            "officedocument.spreadsheetml.sheet"
        ),
        headers={
            "Content-Disposition": (
                f'attachment; filename="{filename}"'
            )
        }
    )

def parse_excel_date(value):
    if not value:
        return None

    try:
        return date.fromisoformat(value)
    except (TypeError, ValueError):
        return value


def style_export_sheet(
    worksheet,
    column_widths: list[int]
):
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    worksheet.sheet_view.showGridLines = False

    header_fill = PatternFill(
        fill_type="solid",
        fgColor="2563EB"
    )

    for cell in worksheet[1]:
        cell.font = Font(
            bold=True,
            color="FFFFFF"
        )
        cell.fill = header_fill
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True
        )

    worksheet.row_dimensions[1].height = 28

    for column_index, width in enumerate(
        column_widths,
        start=1
    ):
        column_letter = get_column_letter(
            column_index
        )
        worksheet.column_dimensions[
            column_letter
        ].width = width

    for row in worksheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(
                vertical="top",
                wrap_text=True
            )
